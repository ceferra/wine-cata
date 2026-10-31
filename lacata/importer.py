"""Import wines from a CSV or Excel file (e.g. a template filled in by ChatGPT).

The template has one row per wine. Headers are accepted in Spanish, English or Valencian (and
the internal field names), values are normalised to the app's canonical options where possible
and anything doubtful is reported as a warning instead of being silently dropped.
"""

import csv
import io
import re
import unicodedata

import pandas as pd

from .options import DEFAULT_AGINGS, DEFAULT_GRAPES, ORIGIN_REGIONS, WINE_TYPES

# field -> template header per language (es, en, va)
COLUMNS = {
    "name": ("Nombre", "Name", "Nom"),
    "winery": ("Bodega", "Winery", "Celler"),
    "grape": ("Uva", "Grape", "Raïm"),
    "origin": ("D.O.", "Region", "D.O."),
    "aging": ("Crianza", "Aging", "Criança"),
    "type": ("Tipo", "Type", "Tipus"),
    "year": ("Añada", "Vintage", "Anyada"),
    "alcohol": ("Alcohol", "Alcohol", "Alcohol"),
    "price": ("Precio", "Price", "Preu"),
    "rating": ("Puntuación", "Rating", "Puntuació"),
    "description": ("Descripción", "Description", "Descripció"),
}
_LANG_IDX = {"es": 0, "en": 1, "va": 2}

EXAMPLE = {"name": "Protos Roble", "winery": "Protos", "grape": "Tempranillo", "origin": "Ribera del Duero",
           "aging": "Roble", "type": "Tinto", "year": "2022", "alcohol": 14.0, "price": 11.0, "rating": 3.8,
           "description": "Fruta roja madura y un toque tostado."}

_TYPE_ALIASES = {
    "tinto": "Tinto", "red": "Tinto", "negre": "Tinto",
    "blanco": "Blanco", "white": "Blanco", "blanc": "Blanco",
    "rosado": "Rosado", "rose": "Rosado", "rosat": "Rosado",
    "espumoso": "Espumoso", "sparkling": "Espumoso", "escumos": "Espumoso", "cava": "Espumoso",
    "dulce": "Dulce", "sweet": "Dulce", "dolc": "Dulce",
    "generoso": "Generoso", "fortified": "Generoso", "generos": "Generoso",
}
_AGING_ALIASES = {"young": "Joven", "jove": "Joven", "oak": "Roble", "roure": "Roble", "crianca": "Crianza",
                  "reserve": "Reserva", "gran reserve": "Gran Reserva"}


class ImportFailed(Exception):
    """The file could not be read at all (bad format, no name column)."""


def _key(s) -> str:
    """Lower-case, accent-free, single-spaced version of s for matching."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


_HEADER_MAP = {_key(h): f for f, hs in COLUMNS.items() for h in (*hs, f)}
_HEADER_MAP.update({"do": "origin", "denominacion de origen": "origin", "origen": "origin", "origin": "origin",
                    "uvas": "grape", "grapes": "grape", "variedad": "grape", "variedades": "grape",
                    "ano": "year", "year": "year", "puntuacion": "rating", "precio (eur)": "price",
                    "alcohol (%)": "alcohol", "vino": "name", "wine": "name"})


def headers(lang="es") -> list:
    return [hs[_LANG_IDX.get(lang, 0)] for hs in COLUMNS.values()]


# ---------------------------------------------------------------- templates

def template_csv(lang="es") -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers(lang))
    w.writerow([EXAMPLE[f] for f in COLUMNS])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens accents correctly


def template_xlsx(opts, lang="es") -> bytes:
    """Excel template: 'Vinos' sheet with an example row and drop-downs, plus a sheet of valid values."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Vinos"
    ws.append(headers(lang))
    ws.append([EXAMPLE[f] for f in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7B1E3A")
    widths = {"name": 32, "winery": 20, "grape": 28, "origin": 22, "description": 45}
    for i, f in enumerate(COLUMNS):
        ws.column_dimensions[chr(65 + i)].width = widths.get(f, 13)
    ws.freeze_panes = "A2"

    lists = {"aging": opts["agings"], "type": WINE_TYPES, "origin": opts["origins"], "grape": opts["grapes"]}
    vs = wb.create_sheet("Valores")
    for col, (f, values) in enumerate(lists.items(), start=1):
        letter = chr(64 + col)
        vs.cell(1, col, headers(lang)[list(COLUMNS).index(f)]).font = Font(bold=True)
        for row, v in enumerate(values, start=2):
            vs.cell(row, col, v)
        vs.column_dimensions[letter].width = 26
        if f in ("aging", "type", "origin"):  # grapes can be several per cell, so no drop-down
            dv = DataValidation(type="list", formula1=f"=Valores!${letter}$2:${letter}${len(values) + 1}",
                                allow_blank=True, showErrorMessage=f != "origin")  # other D.O.s are allowed
            ws.add_data_validation(dv)
            dv.add(f"{chr(65 + list(COLUMNS).index(f))}2:{chr(65 + list(COLUMNS).index(f))}500")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def chatgpt_prompt(opts, lang="es") -> str:
    h = headers(lang)
    return (
        "Rellena una tabla de vinos para una cata a ciegas. Devuélvela como archivo Excel (.xlsx) o como CSV, "
        "una fila por vino, con exactamente estas columnas y en este orden:\n"
        f"{', '.join(h)}\n\n"
        "Reglas:\n"
        f"- {h[2]}: una o varias uvas separadas por comas. Usa estos nombres cuando sea posible: "
        f"{', '.join(opts['grapes'])}.\n"
        f"- {h[3]}: la denominación de origen, sin «D.O.» delante. Usa una de estas si corresponde: "
        f"{', '.join(opts['origins'])}.\n"
        f"- {h[4]}: una de {', '.join(opts['agings'])} (vacío si no aplica, p. ej. espumosos o generosos).\n"
        f"- {h[5]}: uno de {', '.join(WINE_TYPES)}.\n"
        f"- {h[6]}: año con 4 cifras, o vacío.\n"
        f"- {h[7]}: grados, número con punto decimal (13.5). {h[8]}: euros por botella en tienda (12.5). "
        f"{h[9]}: nota media de 0 a 5 tipo Vivino (3.8).\n"
        f"- {h[10]}: una o dos frases de nota de cata.\n"
        "- Si no estás seguro de un dato, déjalo vacío en lugar de inventarlo.\n\n"
        "Los vinos son:\n"
    )


# ---------------------------------------------------------------- parsing

def read_table(data: bytes, filename: str) -> pd.DataFrame:
    name = (filename or "").lower()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            return pd.read_excel(io.BytesIO(data), sheet_name=0, dtype=str)
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return parse_text_table(text)
    except ImportFailed:
        raise
    except Exception as e:  # pandas/openpyxl raise many different exception types for bad files
        raise ImportFailed(str(e)) from e


def parse_text_table(text: str) -> pd.DataFrame:
    """CSV (comma, semicolon or tab separated) or a Markdown table as pasted from ChatGPT."""
    lines = [ln.strip() for ln in text.strip().splitlines() if ln.strip()]
    lines = [ln for ln in lines if not ln.startswith("```")]
    if lines and lines[0].startswith("|"):
        rows = [[c.strip() for c in ln.strip("|").split("|")] for ln in lines
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", ln)]
        return pd.DataFrame(rows[1:], columns=rows[0], dtype=str) if rows else pd.DataFrame()
    body = "\n".join(lines)
    try:
        sep = csv.Sniffer().sniff(lines[0], delimiters=",;\t").delimiter if lines else ","
    except csv.Error:
        sep = ","
    return pd.read_csv(io.StringIO(body), sep=sep, dtype=str, keep_default_na=False)


def _num(v):
    if v in (None, ""):
        return None
    s = str(v).replace("€", "").replace("%", "").replace("º", "").replace("°", "").strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+,\d+", s):  # 1.234,50
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return "bad"


def _match(value, options, aliases=None):
    """Canonical spelling of value among options (case/accents ignored), or None."""
    k = _key(value)
    for o in options:
        if _key(o) == k:
            return o
    if aliases and k in aliases:
        return aliases[k]
    return None


def rows_to_wines(df: pd.DataFrame, opts) -> tuple:
    """Returns (wines, warnings, ignored_columns). Rows without a name are skipped."""
    cols = {}
    ignored = []
    for c in df.columns:
        f = _HEADER_MAP.get(_key(c))
        if f and f not in cols.values():
            cols[c] = f
        elif str(c).strip() and not str(c).startswith("Unnamed"):
            ignored.append(str(c))
    if "name" not in cols.values():
        raise ImportFailed("no_name_column")

    known_origins = list(dict.fromkeys([*opts["origins"], *ORIGIN_REGIONS]))
    known_grapes = list(dict.fromkeys([*opts["grapes"], *DEFAULT_GRAPES]))
    wines, warnings, seen = [], [], set()
    for i, raw in enumerate(df.fillna("").to_dict("records"), start=2):  # row 1 is the header
        r = {f: str(raw[c]).strip() for c, f in cols.items()}
        r = {f: ("" if v.lower() in ("nan", "none", "-", "n/a") else v) for f, v in r.items()}
        if not r.get("name"):
            continue
        w = {f: r.get(f, "") for f in COLUMNS}
        w["source"] = "manual"
        label = f"{i} ({w['name']})"
        if _key(w["name"]) in seen:
            warnings.append(("dup", label, ""))
            continue
        seen.add(_key(w["name"]))

        grapes = [g.strip() for g in re.split(r"[,;/]|\s+y\s+|\s+and\s+|\s+i\s+", w["grape"]) if g.strip()]
        w["grape"] = ", ".join(_match(g, known_grapes) or g.title() for g in grapes)

        origin = re.sub(r"^(d\.?\s?o\.?(\s?(ca|q|p)\.?)?|denominaci[oó]n de origen( calificada)?)(?=[\s.]|$)\s*",
                        "", w["origin"], flags=re.I).strip()
        if origin:
            w["origin"] = _match(origin, known_origins) or origin
            if w["origin"] not in ORIGIN_REGIONS:
                warnings.append(("origin", label, w["origin"]))
        if w["aging"]:
            a = _match(w["aging"], opts["agings"], _AGING_ALIASES)
            if a is None:
                warnings.append(("aging", label, w["aging"]))
            w["aging"] = a or ""
        if w["type"]:
            tp = _match(w["type"], WINE_TYPES, _TYPE_ALIASES)
            if tp is None:
                warnings.append(("type", label, w["type"]))
            w["type"] = tp or ""
        m = re.search(r"\b(19|20)\d{2}\b", w["year"])
        w["year"] = m.group(0) if m else ""
        for f, lo, hi in (("alcohol", 0, 25), ("price", 0, 5000), ("rating", 0, 5)):
            v = _num(w[f])
            if v == "bad" or (v is not None and not lo <= v <= hi):
                warnings.append(("number", label, f"{COLUMNS[f][0]}: {w[f]}"))
                v = None
            w[f] = v
        wines.append(w)
    return wines, warnings, ignored
